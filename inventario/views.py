# Protege las páginas del sistema
from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
# Funciones para iniciar y cerrar sesión
from django.contrib.auth import authenticate, login, logout
from openpyxl import load_workbook
from django.db.models import Q
from django.db import transaction
from django.core.paginator import Paginator

from django.utils import timezone

# Importamos los modelos
from .models import Categoria, Ubicacion, Custodio, Bien, Asignacion, Constatacion, Baja 

# Importamos los formularios
from .forms import CategoriaForm, UbicacionForm, CustodioForm, BienForm, AsignacionForm, ConstatacionForm, BajaForm ,ImportarBienesForm, ImportarCustodiosForm

# Muestra el resumen general del sistema
@login_required
def dashboard(request):

    # Total de bienes registrados
    total_bienes = Bien.objects.count()

    # IDs de los bienes que están actualmente asignados
    bienes_ocupados = Asignacion.objects.filter(
        fecha_devolucion__isnull=True,
        bienes__isnull=False
    ).values_list(
        'bienes__id',
        flat=True
    ).distinct()

    # Cantidad real de bienes asignados
    bienes_asignados = Bien.objects.filter(
        id__in=bienes_ocupados
    ).count()

    # Bienes activos que no están asignados
    disponibles = Bien.objects.filter(
        activo=True
    ).exclude(
        id__in=bienes_ocupados
    ).count()

    # Bienes dados de baja
    dados_baja = Baja.objects.values(
        'bien'
    ).distinct().count()

    contexto = {
        'total_bienes': total_bienes,
        'bienes_asignados': bienes_asignados,
        'disponibles': disponibles,
        'dados_baja': dados_baja,
    }

    return render(
        request,
        'inventario/dashboard.html',
        contexto
    )
# Muestra la lista de categorías
@login_required
def lista_categorias(request):
    categorias = Categoria.objects.all()

    return render(
        request,
        'inventario/categorias/lista.html',
        {'categorias': categorias}
    )

# Registra una nueva categoría
@login_required
def nueva_categoria(request):

    if request.method == 'POST':
        form = CategoriaForm(request.POST)

        if form.is_valid():
            form.save()
            return redirect('lista_categorias')

    else:
        form = CategoriaForm()

    return render(
        request,
        'inventario/categorias/formulario.html',
        {'form': form}
    )

# Edita una categoría existente
@login_required
def editar_categoria(request, id):

    categoria = Categoria.objects.get(id=id)

    if request.method == 'POST':
        form = CategoriaForm(request.POST, instance=categoria)

        if form.is_valid():
            form.save()
            return redirect('lista_categorias')

    else:
        form = CategoriaForm(instance=categoria)

    return render(
        request,
        'inventario/categorias/formulario.html',
        {'form': form}
    )


# Activa o desactiva una categoría
@login_required
def cambiar_estado_categoria(request, id):

    categoria = Categoria.objects.get(id=id)

    categoria.activo = not categoria.activo
    categoria.save()

    return redirect('lista_categorias')

# Muestra la lista de ubicaciones
@login_required
def lista_ubicaciones(request):
    ubicaciones = Ubicacion.objects.all()

    return render(
        request,
        'inventario/ubicaciones/lista.html',
        {'ubicaciones': ubicaciones}
    )


# Registra una nueva ubicación
@login_required
def nueva_ubicacion(request):

    if request.method == 'POST':
        form = UbicacionForm(request.POST)

        if form.is_valid():
            form.save()
            return redirect('lista_ubicaciones')

    else:
        form = UbicacionForm()

    return render(
        request,
        'inventario/ubicaciones/formulario.html',
        {'form': form}
    )

# Edita una ubicación
@login_required
def editar_ubicacion(request, id):

    ubicacion = Ubicacion.objects.get(id=id)

    if request.method == 'POST':
        form = UbicacionForm(request.POST, instance=ubicacion)

        if form.is_valid():
            form.save()
            return redirect('lista_ubicaciones')

    else:
        form = UbicacionForm(instance=ubicacion)

    return render(
        request,
        'inventario/ubicaciones/formulario.html',
        {'form': form}
    )


# Activa o desactiva una ubicación
@login_required
def cambiar_estado_ubicacion(request, id):
    ubicacion = Ubicacion.objects.get(id=id)

    ubicacion.activo = not ubicacion.activo
    ubicacion.save()

    return redirect('lista_ubicaciones')


# Muestra la lista de custodios
@login_required
def lista_custodios(request):
    custodios = Custodio.objects.all()

    return render(
        request,
        'inventario/custodios/lista.html',
        {'custodios': custodios}
    )


# Registra un nuevo custodio
@login_required
def nuevo_custodio(request):

    if request.method == 'POST':
        form = CustodioForm(request.POST)

        if form.is_valid():
            form.save()
            return redirect('lista_custodios')

    else:
        form = CustodioForm()

    return render(
        request,
        'inventario/custodios/formulario.html',
        {'form': form}
    )

# Edita un custodio
@login_required
def editar_custodio(request, id):
    custodio = Custodio.objects.get(id=id)

    if request.method == 'POST':
        form = CustodioForm(request.POST, instance=custodio)

        if form.is_valid():
            form.save()
            return redirect('lista_custodios')
    else:
        form = CustodioForm(instance=custodio)

    return render(
        request,
        'inventario/custodios/formulario.html',
        {'form': form}
    )

# Activa o desactiva un custodio
@login_required
def cambiar_estado_custodio(request, id):
    custodio = Custodio.objects.get(id=id)

    custodio.activo = not custodio.activo
    custodio.save()

    return redirect('lista_custodios')

# Muestra la lista de bienes registrados
@login_required
def lista_bienes(request):

    # Todos los bienes
    bienes_queryset = Bien.objects.all().order_by('id')

    # Texto del buscador
    buscar = request.GET.get('buscar', '').strip()

    # Aplica búsqueda
    if buscar:
        bienes_queryset = bienes_queryset.filter(
            Q(codigo__icontains=buscar) |
            Q(nombre__icontains=buscar) |
            Q(marca__icontains=buscar) |
            Q(modelo__icontains=buscar) |
            Q(serie__icontains=buscar)
        )

    # Todos los bienes para imprimir
    bienes_reporte = bienes_queryset

    # Solo 20 bienes por página para la pantalla
    paginador = Paginator(bienes_queryset, 20)
    numero_pagina = request.GET.get('page')
    bienes_pagina = paginador.get_page(numero_pagina)

    return render(
        request,
        'inventario/bienes/lista.html',
        {
            'bienes': bienes_pagina,
            'bienes_reporte': bienes_reporte,
            'buscar': buscar,
        }
    )
# Muestra toda la información de un bien
@login_required
def detalle_bien(request, id):
    bien = get_object_or_404(Bien, id=id)

    return render(
        request,
        'inventario/bienes/detalle.html',
        {'bien': bien}
    )

    

# Registra un nuevo bien
@login_required
def nuevo_bien(request):

    if request.method == 'POST':

        # request.FILES permite guardar la fotografía
        form = BienForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():
            form.save()
            return redirect('lista_bienes')

    else:
        form = BienForm()

    return render(
        request,
        'inventario/bienes/formulario.html',
        {'form': form}
    )
# Edita un bien registrado
@login_required
def editar_bien(request, id):

    bien = get_object_or_404(Bien, id=id)

    if request.method == 'POST':

        # Permite cambiar o conservar la fotografía
        form = BienForm(
            request.POST,
            request.FILES,
            instance=bien
        )

        if form.is_valid():
            form.save()
            return redirect('lista_bienes')

    else:
        form = BienForm(instance=bien)

    return render(
        request,
        'inventario/bienes/formulario.html',
        {
            'form': form,
            'bien': bien
        }
    )

# Activa o desactiva un bien
@login_required
def cambiar_estado_bien(request, id):
    bien = Bien.objects.get(id=id)

    bien.activo = not bien.activo
    bien.save()

    return redirect('lista_bienes')

# Muestra la lista de asignaciones
@login_required
def lista_asignaciones(request):

    asignaciones = Asignacion.objects.all()

    # Texto buscado
    buscar = request.GET.get('buscar', '').strip()

    # Estado seleccionado
    estado = request.GET.get('estado', '').strip()

    # Busca dentro de todos los bienes de la asignación
    if buscar:
        asignaciones = asignaciones.filter(
            Q(bienes__nombre__icontains=buscar) |
            Q(bienes__codigo__icontains=buscar) |
            Q(custodio__nombres__icontains=buscar) |
            Q(custodio__apellidos__icontains=buscar) |
            Q(numero_acta__icontains=buscar)
        ).distinct()

    # Asignaciones pendientes
    if estado == 'pendiente':
        asignaciones = asignaciones.filter(
            fecha_devolucion__isnull=True
        )

    # Asignaciones devueltas
    elif estado == 'devuelto':
        asignaciones = asignaciones.filter(
            fecha_devolucion__isnull=False
        )

    return render(
        request,
        'inventario/asignaciones/lista.html',
        {
            'asignaciones': asignaciones,
            'buscar': buscar,
            'estado': estado
        }
    )

# Registra una nueva asignación
@login_required
def nueva_asignacion(request):

   # Bienes que están en asignaciones todavía no devueltas
    bienes_ocupados = Asignacion.objects.filter(
        fecha_devolucion__isnull=True,
        bienes__isnull=False
    ).values_list(
        'bienes__id',
        flat=True
    ).distinct()

    bienes_disponibles = Bien.objects.filter(
        activo=True
    ).exclude(
        id__in=bienes_ocupados
    )

    if request.method == 'POST':

        form = AsignacionForm(request.POST)

        # Recibe los bienes seleccionados
        bienes_ids = request.POST.getlist('bienes_seleccionados')
        
        if form.is_valid() and bienes_ids:

            # Comprueba que los bienes sigan disponibles
            bienes = Bien.objects.filter(
                id__in=bienes_ids,
                activo=True
            ).exclude(
                id__in=bienes_ocupados
            )

            # Todos los seleccionados deben estar disponibles
            if bienes.count() == len(set(bienes_ids)):

                asignacion = form.save()

                # Guarda todos los bienes seleccionados
                asignacion.bienes.set(bienes)

                return redirect('lista_asignaciones')

        else:
            bienes_ids = []

    else:
        form = AsignacionForm()
        bienes_ids = []

    return render(
        request,
        'inventario/asignaciones/formulario.html',
        {
            'form': form,
            'bienes_disponibles': bienes_disponibles,
            'bienes_ids': bienes_ids
        }
    )

# Edita una asignación
@login_required
def editar_asignacion(request, id):

    asignacion = get_object_or_404(Asignacion, id=id)

    # Bienes que ya pertenecen a esta asignación
    bienes_actuales = asignacion.bienes.all()

    # Compatibilidad con asignaciones antiguas
    if not bienes_actuales.exists() and asignacion.bien:
        bienes_actuales = Bien.objects.filter(
            id=asignacion.bien.id
        )

    # Bienes ocupados por otras asignaciones pendientes
    bienes_ocupados = Asignacion.objects.filter(
        fecha_devolucion__isnull=True,
        bienes__isnull=False
    ).exclude(
        id=asignacion.id
    ).values_list(
        'bienes__id',
        flat=True
    ).distinct()

    # Bienes disponibles para agregar
    bienes_disponibles = Bien.objects.filter(
        activo=True
    ).exclude(
        id__in=bienes_ocupados
    ).exclude(
        id__in=bienes_actuales.values_list('id', flat=True)
    )

    if request.method == 'POST':

        form = AsignacionForm(
            request.POST,
            instance=asignacion
        )

        # Bienes seleccionados en el formulario
        bienes_ids = request.POST.getlist(
            'bienes_seleccionados'
        )

        if form.is_valid() and bienes_ids:

            # Comprueba que no estén ocupados por otra asignación
            bienes = Bien.objects.filter(
                id__in=bienes_ids,
                activo=True
            ).exclude(
                id__in=bienes_ocupados
            )

            if bienes.count() == len(set(bienes_ids)):

                asignacion = form.save()

                # Actualiza todos los bienes
                asignacion.bienes.set(bienes)

                return redirect('lista_asignaciones')

    else:
        form = AsignacionForm(
            instance=asignacion
        )

    return render(
        request,
        'inventario/asignaciones/formulario.html',
        {
            'form': form,
            'bienes_disponibles': bienes_disponibles,
            'bienes_actuales': bienes_actuales,
        }
    )
# Registra la devolución de un bien
@login_required
def registrar_devolucion(request, id):
    asignacion = Asignacion.objects.get(id=id)

    asignacion.fecha_devolucion = timezone.localdate()
    asignacion.save()

    return redirect('lista_asignaciones')

# Muestra el acta de entrega-recepción
@login_required
def acta_asignacion(request, id):

    asignacion = Asignacion.objects.get(id=id)

    return render(
        request,
        'inventario/asignaciones/acta.html',
        {'asignacion': asignacion}
    )

# Elimina una asignación
@login_required
def eliminar_asignacion(request, id):

    # Busca la asignación
    asignacion = get_object_or_404(Asignacion, id=id)

    # Solo elimina si se envía desde el botón Eliminar
    if request.method == 'POST':
        asignacion.delete()

    # Regresa a la lista de asignaciones
    return redirect('lista_asignaciones')

# Muestra la lista de constataciones
@login_required
def lista_constataciones(request):
    constataciones = Constatacion.objects.all()

    return render(
        request,
        'inventario/constataciones/lista.html',
        {'constataciones': constataciones}
    
    )

# Registra una nueva constatación
@login_required
def nueva_constatacion(request):

    if request.method == 'POST':
        form = ConstatacionForm(request.POST)

        if form.is_valid():
            form.save()
            return redirect('lista_constataciones')

    else:
        form = ConstatacionForm()

    return render(
        request,
        'inventario/constataciones/formulario.html',
        {'form': form}
    )

# Edita una constatación
@login_required
def editar_constatacion(request, id):
    constatacion = Constatacion.objects.get(id=id)

    if request.method == 'POST':
        form = ConstatacionForm(
            request.POST,
            instance=constatacion
        )

        if form.is_valid():
            form.save()
            return redirect('lista_constataciones')
    else:
        form = ConstatacionForm(instance=constatacion)

    return render(
        request,
        'inventario/constataciones/formulario.html',
        {'form': form}
    )

# Muestra la lista de bajas
@login_required
def lista_bajas(request):
    bajas = Baja.objects.all()

    return render(
        request,
        'inventario/bajas/lista.html',
        {'bajas': bajas}
    )

# Registra una nueva baja
@login_required
def nueva_baja(request):

    if request.method == 'POST':
        form = BajaForm(request.POST)

        if form.is_valid():
            baja = form.save()

            # Desactiva automáticamente el bien
            baja.bien.activo = False
            baja.bien.save()

            return redirect('lista_bajas')

    else:
        form = BajaForm()

    return render(
        request,
        'inventario/bajas/formulario.html',
        {'form': form}
    )

# Muestra la página de reportes
@login_required
def reportes(request):

    return render(
        request,
        'inventario/reportes.html'
    )

# Inicio de sesión del administrador

def iniciar_sesion(request):

    mensaje = ''

    if request.method == 'POST':
        usuario = request.POST.get('usuario')
        clave = request.POST.get('clave')

        user = authenticate(
            request,
            username=usuario,
            password=clave
        )

        if user is not None:
            login(request, user)
            return redirect('dashboard')
        else:
            mensaje = 'Usuario o contraseña incorrectos.'

    return render(
        request,
        'inventario/login.html',
        {'mensaje': mensaje}
    )


# Cierra la sesión
@login_required
def cerrar_sesion(request):
    logout(request)
    return redirect('login')

# Importa bienes desde Excel conservando los datos originales del ISTAM
@login_required
def importar_bienes(request):

    mensaje = ''

    if request.method == 'POST':
        form = ImportarBienesForm(request.POST, request.FILES)

        if form.is_valid():
            archivo = request.FILES['archivo']
            # Procedencia seleccionada
            procedencia = form.cleaned_data['procedencia']

            try:
                libro = load_workbook(archivo, data_only=True)
                hoja = libro.active

                creados = 0
                duplicados = 0
                errores = 0
                filas_error = []

                categoria, _ = Categoria.objects.get_or_create(
                    nombre='Sin categoría'
                )

                ubicacion, _ = Ubicacion.objects.get_or_create(
                    nombre='Por definir'
                )

                def texto_original(valor):
                    if valor is None:
                        return None

                    return str(valor)

                def numero(valor):
                    if valor is None or valor == '':
                        return None

                    try:
                        if isinstance(valor, str):
                            valor = valor.replace(',', '.').strip()

                        return float(valor)

                    except (ValueError, TypeError):
                        return None

                # Los datos empiezan en la fila 6
                for numero_fila, fila in enumerate(
                    hoja.iter_rows(min_row=6, values_only=True),
                    start=6
                ):

                    try:
                        # Columnas principales que ya usábamos
                        codigo = fila[0]
                        tipo_bien = fila[1]
                        nombre = fila[2]
                        caracteristicas = fila[3]
                        color = fila[4]
                        alto = fila[5]
                        ancho = fila[6]
                        profundidad = fila[7]
                        estado = fila[8]
                        observaciones = fila[9]

                        # Columnas adicionales originales del ISTAM
                        valor_compra_original = (
                            fila[10] if len(fila) > 10 else None
                        )

                        material = (
                            fila[11] if len(fila) > 11 else None
                        )

                        dimensiones_originales = (
                            fila[12] if len(fila) > 12 else None
                        )

                        ubicacion_original = (
                            fila[13] if len(fila) > 13 else None
                        )

                        custodio_original = (
                            fila[14] if len(fila) > 14 else None
                        )

                        serie = (
                            fila[15] if len(fila) > 15 else None
                        )

                        modelo = (
                            fila[16] if len(fila) > 16 else None
                        )

                        marca = (
                            fila[17] if len(fila) > 17 else None
                        )
                        datos_adicionales_originales = (
                            fila[18] if len(fila) > 18 else None
                        )

                        # Ignora filas completamente vacías
                        if all(valor is None for valor in fila):
                            continue

                        
                        # Si tiene código pero no descripción, sí es error
                        if nombre is None or str(nombre).strip() == '':
                            errores += 1
                            filas_error.append(numero_fila)
                            continue

                        # El código puede quedar pendiente
                        if codigo is not None and str(codigo).strip() != '':
                            codigo = str(codigo).strip()
                        else:
                            codigo = None
                        nombre = str(nombre).strip()

                        # Evita duplicados por código
                        if codigo and Bien.objects.filter(codigo=codigo).exists():
                            duplicados += 1
                            continue

                        # Conserva el estado original
                        condicion_original = texto_original(estado)

                        # Para el campo interno del sistema:
                        # solo usa valores válidos
                        estado_sistema = 'BUENO'

                        if estado:
                            estado_mayuscula = str(estado).strip().upper()

                            if estado_mayuscula in [
                                'BUENO',
                                'REGULAR',
                                'MALO'
                            ]:
                                estado_sistema = estado_mayuscula

                        Bien.objects.create(
                            codigo=codigo,
                            nombre=nombre,
                            categoria=categoria,
                            ubicacion=ubicacion,

                            tipo_bien=texto_original(tipo_bien),
                            marca=texto_original(marca),
                            modelo=texto_original(modelo),
                            serie=texto_original(serie),

                            valor_compra=numero(valor_compra_original),
                            valor_compra_original=texto_original(
                                valor_compra_original
                            ),

                            color=texto_original(color),
                            material=texto_original(material),
                            caracteristicas=texto_original(
                                caracteristicas
                            ),

                            alto=numero(alto),
                            ancho=numero(ancho),
                            profundidad=numero(profundidad),

                            dimensiones_originales=texto_original(
                                dimensiones_originales
                            ),

                            estado=estado_sistema,
                            condicion_original=condicion_original,

                            ubicacion_original=texto_original(
                                ubicacion_original
                            ),

                            custodio_original=texto_original(
                                custodio_original
                            ),

                            observaciones=texto_original(
                                observaciones
                            ),
                            datos_adicionales_originales=texto_original(
                                datos_adicionales_originales
                            ),

                            procedencia=procedencia,

                            activo=True
                        )

                        creados += 1

                    except Exception:
                        errores += 1
                        filas_error.append(numero_fila)
                        continue

                mensaje = (
                    f'Importación finalizada. '
                    f'Registrados: {creados}. '
                    f'Duplicados: {duplicados}. '
                    f'Con errores: {errores}.'
                )

                if filas_error:
                    mensaje += (
                        ' Revisar filas: '
                        + ', '.join(map(str, filas_error))
                    )

            except Exception as error:
                mensaje = (
                    f'No se pudo leer el archivo Excel. '
                    f'Error: {error}'
                )

    else:
        form = ImportarBienesForm()

    return render(
        request,
        'inventario/bienes/importar.html',
        {
            'form': form,
            'mensaje': mensaje
        }
    )

# Importa custodios desde un archivo Excel
@login_required
def importar_custodios(request):

    mensaje = ''

    if request.method == 'POST':

        form = ImportarCustodiosForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            archivo = request.FILES['archivo']

            try:
                libro = load_workbook(
                    archivo,
                    data_only=True
                )

                hoja = libro.active

                creados = 0
                duplicados = 0
                errores = 0

                # Los datos empiezan en la fila 2
                for fila in hoja.iter_rows(
                    min_row=2,
                    values_only=True
                ):

                    try:

                        cedula = fila[0]
                        nombres = fila[1]
                        apellidos = fila[2]
                        cargo = fila[3]
                        departamento = fila[4]
                        correo = fila[5]
                        telefono = fila[6]

                        # Ignora filas vacías
                        if all(valor is None for valor in fila):
                            continue

                        # Datos obligatorios
                        if not cedula or not nombres or not apellidos:
                            errores += 1
                            continue

                        cedula = str(cedula).strip()
                        nombres = str(nombres).strip()
                        apellidos = str(apellidos).strip()

                        # Evita cédulas repetidas
                        if Custodio.objects.filter(
                            cedula=cedula
                        ).exists():

                            duplicados += 1
                            continue

                        Custodio.objects.create(
                            cedula=cedula,
                            nombres=nombres,
                            apellidos=apellidos,
                            cargo=str(cargo).strip()
                            if cargo else '',

                            departamento=str(
                                departamento
                            ).strip()
                            if departamento else '',

                            correo=str(correo).strip()
                            if correo else '',

                            telefono=str(telefono).strip()
                            if telefono else '',

                            activo=True
                        )

                        creados += 1

                    except Exception:
                        errores += 1
                        continue

                mensaje = (
                    f'Importación finalizada. '
                    f'Registrados: {creados}. '
                    f'Duplicados: {duplicados}. '
                    f'Con errores: {errores}.'
                )

            except Exception as error:

                mensaje = (
                    'No se pudo leer el archivo Excel. '
                    f'Error: {error}'
                )

    else:
        form = ImportarCustodiosForm()

    return render(
        request,
        'inventario/custodios/importar.html',
        {
            'form': form,
            'mensaje': mensaje
        }
    )


# Importa la matriz de entrega de actas
@login_required
def importar_matriz_actas(request):

    mensaje = ''

    if request.method == 'POST':

        archivo = request.FILES.get('archivo')

        if not archivo:
            mensaje = 'Seleccione un archivo Excel.'

        else:
            try:
                # Abre el archivo Excel
                libro = load_workbook(
                    archivo,
                    data_only=True
                )

                hoja = libro.active

                # Categoría de los portátiles
                categoria, _ = Categoria.objects.get_or_create(
                    nombre='Computadoras'
                )

                # Ubicación temporal
                ubicacion, _ = Ubicacion.objects.get_or_create(
                    nombre='Por definir'
                )

                creados = 0
                duplicados = 0
                errores = 0
                detalle_errores = []

                # Convierte valores del Excel a texto
                def texto(valor):
                    if valor is None:
                        return ''
                    return str(valor).strip()

                # Recorre las filas del Excel
                for numero_fila, fila in enumerate(
                    hoja.iter_rows(
                        min_row=2,
                        values_only=True
                    ),
                    start=2
                ):

                    try:
                        # Ignora filas completamente vacías
                        if all(valor is None for valor in fila):
                            continue

                        # -------------------------
                        # DATOS DEL CUSTODIO
                        # -------------------------

                        numero_acta = texto(fila[0])
                        nombre_completo = texto(fila[1])
                        cedula = texto(fila[2])
                        celular = texto(fila[3])
                        ciclo = texto(fila[4])
                        carrera = texto(fila[5])
                        direccion = texto(fila[6])

                        # -------------------------
                        # DATOS DEL PORTÁTIL
                        # -------------------------

                        marca = texto(fila[7])
                        modelo = texto(fila[8])
                        serie = texto(fila[9])
                        codigo = texto(fila[10])

                        # -------------------------
                        # DATOS DEL CARGADOR
                        # -------------------------

                        cargador_marca = texto(fila[11])
                        cargador_modelo = texto(fila[12])
                        cargador_serie = texto(fila[13])

                        # -------------------------
                        # DATOS DEL MOUSE
                        # -------------------------

                        mouse_marca = texto(fila[14])
                        mouse_modelo = texto(fila[15])
                        mouse_serie = texto(fila[16])

                        # -------------------------
                        # DATOS DE LA MOCHILA
                        # -------------------------

                        mochila_marca = texto(fila[17])
                        mochila_modelo = texto(fila[18])
                        mochila_serie = texto(fila[19])

                        # Observación
                        observacion = texto(fila[20])

                        # Ignora filas auxiliares que no tienen persona
                        if not cedula or not nombre_completo:
                            continue

                        # -------------------------
                        # NOMBRES Y APELLIDOS
                        # -------------------------

                        partes = nombre_completo.split()

                        if len(partes) >= 4:
                            apellidos = ' '.join(partes[:2])
                            nombres = ' '.join(partes[2:])

                        elif len(partes) >= 2:
                            apellidos = partes[0]
                            nombres = ' '.join(partes[1:])

                        else:
                            apellidos = ''
                            nombres = nombre_completo

                        # -------------------------
                        # CUSTODIO
                        # -------------------------

                        custodio, _ = Custodio.objects.update_or_create(
                            cedula=cedula,
                            defaults={
                                'nombres': nombres,
                                'apellidos': apellidos,
                                'cargo': 'Estudiante',
                                'departamento': '',
                                'telefono': celular,
                                'ciclo': ciclo,
                                'carrera': carrera,
                                'direccion': direccion,
                                'activo': True
                            }
                        )

                        # -------------------------
                        # EVITA DUPLICAR EL BIEN
                        # -------------------------

                        if codigo and Bien.objects.filter(
                            codigo=codigo
                        ).exists():

                            duplicados += 1
                            continue

                        # -------------------------
                        # ACCESORIOS
                        # -------------------------

                        accesorios = (
                            f'Cargador: {cargador_marca} | '
                            f'{cargador_modelo} | {cargador_serie}\n'

                            f'Mouse: {mouse_marca} | '
                            f'{mouse_modelo} | {mouse_serie}\n'

                            f'Mochila: {mochila_marca} | '
                            f'{mochila_modelo} | {mochila_serie}'
                        )

                        # -------------------------
                        # CREA EL PORTÁTIL
                        # -------------------------

                        bien = Bien.objects.create(
                            codigo=codigo if codigo else None,
                            nombre='PORTÁTIL',
                            categoria=categoria,
                            ubicacion=ubicacion,
                            marca=marca,
                            modelo=modelo,
                            serie=serie,
                            caracteristicas=accesorios,
                            observaciones=observacion,
                            estado='BUENO',
                            activo=True
                        )

                       # Crea la asignación
                        asignacion = Asignacion.objects.create(
                            bien=bien,
                            custodio=custodio,
                            numero_acta=numero_acta,
                            fecha_asignacion=None,
                            observaciones=observacion,
                            activa=True
                        )

                        # Relaciona el bien con la nueva asignación múltiple
                        asignacion.bienes.add(bien)

                        creados += 1

                    except Exception as error:

                        errores += 1

                        detalle_errores.append(
                            f'Fila {numero_fila}: {str(error)}'
                        )

                # Resultado de la importación
                mensaje = (
                    f'Importación finalizada. '
                    f'Registrados: {creados}. '
                    f'Duplicados: {duplicados}. '
                    f'Con errores: {errores}.'
                )

                if detalle_errores:
                    mensaje += (
                        ' Errores: '
                        + ' | '.join(detalle_errores)
                    )

            except Exception as error:

                mensaje = (
                    'No se pudo leer la matriz. '
                    f'Error: {str(error)}'
                )

    return render(
        request,
        'inventario/asignaciones/importar_actas.html',
        {
            'mensaje': mensaje
        }
    )